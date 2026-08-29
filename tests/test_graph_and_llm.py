from __future__ import annotations

from pathlib import Path

from docx import Document

from docx_de_identifier.agents.llm_disambiguate import (
    Classification,
    disambiguate_entities,
)
from docx_de_identifier.config import AppConfig, LlmConfig
from docx_de_identifier.graph import build_graph, route_after_merge
from docx_de_identifier.models import EntityCandidate, EntityType, ResolvedEntity


class FakeEntity:
    text = "Jane Smith"
    label_ = "PERSON"
    start_char = 0
    end_char = 10


class FakeDoc:
    ents = [FakeEntity()]


def fake_nlp(text: str) -> FakeDoc:
    return FakeDoc()


def _resolved_person() -> ResolvedEntity:
    candidate = EntityCandidate("p1", 0, 10, "Jane Smith", EntityType.PERSON, 0.85, "spacy.ner")
    return ResolvedEntity(candidate, "person-01", "[PERSON_01]")


def test_uncertain_or_failed_llm_keeps_redaction() -> None:
    config = AppConfig(llm=LlmConfig(enabled=True, model="local"))
    entity = _resolved_person()

    uncertain = disambiguate_entities(
        [entity], config, lambda _: Classification(is_person=False, confidence=0.6)
    )

    assert uncertain == [entity]


def test_confident_llm_rejection_removes_candidate() -> None:
    config = AppConfig(llm=LlmConfig(enabled=True, model="local", reject_threshold=0.9))

    retained = disambiguate_entities(
        [_resolved_person()],
        config,
        lambda _: Classification(is_person=False, confidence=0.99),
    )

    assert retained == []


def test_route_uses_llm_only_for_eligible_candidates() -> None:
    config = AppConfig(llm=LlmConfig(enabled=True, model="local"))
    state = {"docx_path": "in", "output_path": "out", "resolved_entities": [_resolved_person()]}

    assert route_after_merge(state, config) == "llm"
    assert route_after_merge(state, AppConfig()) == "redact"


def test_compiled_graph_runs_end_to_end_with_injected_local_ner(tmp_path: Path) -> None:
    source = tmp_path / "input.docx"
    output = tmp_path / "output.docx"
    document = Document()
    document.add_paragraph("Jane Smith sent the update")
    document.save(source)

    result = build_graph(AppConfig(), nlp=fake_nlp).invoke(
        {"docx_path": str(source), "output_path": str(output)}
    )

    assert output.exists()
    assert "Jane Smith" not in Document(output).paragraphs[0].text
    assert "[PERSON_01]" in Document(output).paragraphs[0].text
    assert len(result["audit_events"]) == 1
