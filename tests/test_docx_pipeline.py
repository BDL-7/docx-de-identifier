from __future__ import annotations

import base64
from pathlib import Path
from zipfile import ZipFile

from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches

from docx_de_identifier.agents.ingestion import ingestion_agent
from docx_de_identifier.agents.merge import merge_agent
from docx_de_identifier.agents.redact import redact_agent
from docx_de_identifier.agents.regex_detect import regex_detection_agent
from docx_de_identifier.config import AppConfig

PNG_1X1 = base64.b64decode(
    "iVBORw0KGgoAAAANSUhEUgAAAAEAAAABCAQAAAC1HAwCAAAAC0lEQVR42mP8/x8AAusB9Y9ZQmcAAAAASUVORK5CYII="
)


def _add_hyperlink(paragraph: object, text: str, url: str) -> None:
    relationship_id = paragraph.part.relate_to(  # type: ignore[attr-defined]
        url,
        "http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink",
        is_external=True,
    )
    hyperlink = OxmlElement("w:hyperlink")
    hyperlink.set(qn("r:id"), relationship_id)
    run = OxmlElement("w:r")
    value = OxmlElement("w:t")
    value.text = text
    run.append(value)
    hyperlink.append(run)
    paragraph._p.append(hyperlink)  # type: ignore[attr-defined]


def _add_tracked_insertion(paragraph: object, text: str) -> None:
    insertion = OxmlElement("w:ins")
    insertion.set(qn("w:author"), "Identifying Author")
    run = OxmlElement("w:r")
    value = OxmlElement("w:t")
    value.text = text
    run.append(value)
    insertion.append(run)
    paragraph._p.append(insertion)  # type: ignore[attr-defined]


def _create_fixture(path: Path) -> None:
    document = Document()
    paragraph = document.add_paragraph()
    paragraph.add_run("Doe, Jane").bold = True
    paragraph.add_run(" (jdoe) 3/13/2024 10:59 AM contact jane@example.org ")
    _add_hyperlink(paragraph, "Jane Smith profile", "https://example.test/Jane-Smith")
    _add_tracked_insertion(paragraph, " Secret Inserted Name")
    document.add_table(rows=1, cols=1).cell(0, 0).text = "Call lasted 47 minutes 7 seconds"
    document.sections[0].header.paragraphs[0].text = "Private Header"
    document.sections[0].footer.paragraphs[0].text = "Private Footer"
    document.add_comment(paragraph.runs[:1], text="Private Comment", author="Jane Doe")
    image_path = path.with_suffix(".png")
    image_path.write_bytes(PNG_1X1)
    document.add_picture(str(image_path), width=Inches(0.1))
    document.core_properties.author = "Jane Doe"
    document.save(path)


def test_ingestion_redaction_and_surface_removal(tmp_path: Path) -> None:
    source = tmp_path / "input.docx"
    output = tmp_path / "output.docx"
    _create_fixture(source)
    config = AppConfig()

    state: dict[str, object] = {"docx_path": str(source), "output_path": str(output)}
    state.update(ingestion_agent(state))
    state.update(regex_detection_agent(state, config))
    state.update(merge_agent(state, config))
    state.update(redact_agent(state))

    reopened = Document(output)
    body_text = "\n".join(paragraph.text for paragraph in reopened.paragraphs)
    table_text = reopened.tables[0].cell(0, 0).text
    assert "Doe, Jane" not in body_text
    assert "jane@example.org" not in body_text
    assert "Secret Inserted Name" not in body_text
    assert "[PERSON_01]" in body_text
    assert "[EMAIL_01]" in body_text
    assert "[DURATION_01]" in table_text
    assert reopened.core_properties.author == ""

    with ZipFile(output) as archive:
        names = set(archive.namelist())
        document_xml = archive.read("word/document.xml").decode("utf-8")
        relationships = archive.read("word/_rels/document.xml.rels").decode("utf-8")
    assert "word/header1.xml" not in names
    assert "word/footer1.xml" not in names
    assert "word/comments.xml" not in names
    assert not any(name.startswith("word/media/") for name in names)
    assert "<w:ins" not in document_xml
    assert "<w:drawing" not in document_xml
    assert "Jane Smith profile" in document_xml
    assert "https://example.test/Jane-Smith" in relationships
