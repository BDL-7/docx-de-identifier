from __future__ import annotations

import json
from pathlib import Path

import spacy
import yaml
from docx import Document

from docx_de_identifier.cli import main
from docx_de_identifier.config import load_config


def test_config_rejects_hosted_llm_endpoint(tmp_path: Path) -> None:
    config_path = tmp_path / "config.yaml"
    config_path.write_text(
        yaml.safe_dump({"llm": {"endpoint": "https://api.example.test", "enabled": True}}),
        encoding="utf-8",
    )

    try:
        load_config(config_path)
    except ValueError as error:
        assert "loopback" in str(error)
    else:
        raise AssertionError("Hosted LLM endpoint was accepted")


def test_cli_writes_atomic_docx_and_pii_free_audit(tmp_path: Path) -> None:
    model_path = tmp_path / "ner-model"
    nlp = spacy.blank("en")
    ruler = nlp.add_pipe("entity_ruler")
    ruler.add_patterns([{"label": "PERSON", "pattern": "Jane Smith"}])
    nlp.to_disk(model_path)

    source = tmp_path / "source.docx"
    output = tmp_path / "result.docx"
    document = Document()
    document.add_paragraph("Jane Smith emailed jane@example.org 2 hours ago")
    document.save(source)
    config_path = tmp_path / "config.yaml"
    config_path.write_text(yaml.safe_dump({"spacy_model": str(model_path)}), encoding="utf-8")

    exit_code = main([str(source), "--output", str(output), "--config", str(config_path)])

    audit_path = output.with_suffix(".audit.json")
    assert exit_code == 0
    assert output.exists()
    assert audit_path.exists()
    assert "Jane Smith" not in Document(output).paragraphs[0].text
    audit_text = audit_path.read_text(encoding="utf-8")
    assert "Jane Smith" not in audit_text
    assert "jane@example.org" not in audit_text
    audit = json.loads(audit_text)
    assert audit["source_sha256"]
    assert audit["output_sha256"]
    assert len(audit["events"]) == 3
